import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_resume():
    doc = docx.Document()

    # Set standard margins (0.5 inch all around for ATS compact presentation)
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Color Palette
    COLOR_PRIMARY = RGBColor(15, 23, 42)    # Slate 900
    COLOR_ACCENT = RGBColor(37, 99, 235)    # Royal Blue
    COLOR_MUTED = RGBColor(71, 85, 105)     # Slate 600
    COLOR_DARK = RGBColor(30, 41, 59)       # Slate 800

    def add_heading_with_border(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(text.upper())
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_PRIMARY
        
        # Add bottom border in XML
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                         r'<w:bottom w:val="single" w:sz="6" w:space="2" w:color="CBD5E1"/>'
                         r'</w:pBdr>')
        pPr.append(pBdr)

    # 1. Header (Centered)
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_before = Pt(0)
    name_p.paragraph_format.space_after = Pt(1)
    
    name_run = name_p.add_run("AHMAD HASSAN")
    name_run.bold = True
    name_run.font.name = "Arial"
    name_run.font.size = Pt(21)
    name_run.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(3)
    
    sub_run = sub_p.add_run("Computer Science Student · Software Engineer (AI & Automation)")
    sub_run.bold = True
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(10.5)
    sub_run.font.color.rgb = COLOR_ACCENT

    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_before = Pt(0)
    contact_p.paragraph_format.space_after = Pt(6)
    
    contact_text = (
        "Pakistan (Open to Remote / Relocation)  |  "
        "ahmad.hassan.dev@gmail.com  |  "
        "github.com/Ahmad-Haassan  |  "
        "Portfolio: github.com/Ahmad-Haassan/ahmad-portfolio"
    )
    c_run = contact_p.add_run(contact_text)
    c_run.font.name = "Arial"
    c_run.font.size = Pt(9)
    c_run.font.color.rgb = COLOR_MUTED

    # 2. Professional Summary
    add_heading_with_border("Professional Summary")
    summary_p = doc.add_paragraph()
    summary_p.paragraph_format.space_before = Pt(2)
    summary_p.paragraph_format.space_after = Pt(5)
    summary_p.paragraph_format.line_spacing = 1.15
    
    s_run = summary_p.add_run(
        "Computer Science student with rigorous foundations in Data Structures, Algorithms, and Software Engineering "
        "principles. Experienced in building automated systems, privacy-first AI desktop applications, and high-performance "
        "full-stack web platforms using Python, Next.js 16, React 19, and TypeScript. Proven track record of independently "
        "architecting production-grade end-to-end solutions including real-time e-commerce price comparators, offline speech "
        "recognition engines, and enterprise webhook integrations."
    )
    s_run.font.name = "Arial"
    s_run.font.size = Pt(9.5)
    s_run.font.color.rgb = COLOR_DARK

    # 3. Technical Skills
    add_heading_with_border("Technical Skills")
    
    skills = [
        ("Languages:", "Python, C++, TypeScript, JavaScript (ES6+), HTML5, CSS3, SQL"),
        ("Web Frameworks:", "Next.js 16 (App Router), React 19, FastAPI, Node.js, Tailwind CSS, PyQt6, Motion"),
        ("AI & Automation:", "Faster-Whisper (Local STT), Vosk (Offline Wake-Word), Meta WhatsApp Cloud API, Google Calendar API v3, BeautifulSoup4"),
        ("Developer Tools:", "Git, GitHub CLI, Linux / PowerShell, VS Code, REST APIs, Webhook Security, Postman")
    ]
    
    for label, val in skills:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        lbl_run = p.add_run(label + " ")
        lbl_run.bold = True
        lbl_run.font.name = "Arial"
        lbl_run.font.size = Pt(9.5)
        lbl_run.font.color.rgb = COLOR_PRIMARY
        
        val_run = p.add_run(val)
        val_run.font.name = "Arial"
        val_run.font.size = Pt(9.5)
        val_run.font.color.rgb = COLOR_DARK

    # 4. Key Engineering Projects
    add_heading_with_border("Key Engineering Projects")

    projects = [
        {
            "title": "ShopLens AI — Multi-Store Price Comparison & Review Sentiment Engine",
            "tech": "Next.js 14, React, TypeScript, Python, BeautifulSoup4, Tailwind CSS",
            "link": "github.com/Ahmad-Haassan/shoplens",
            "bullets": [
                ("Cross-Platform Scraping Architecture: ", "Built concurrent data extraction scrapers pulling real-time pricing, stock availability, and shipping metrics across Amazon, Best Buy, and Walmart."),
                ("Normalization & Matching Engine: ", "Engineered a variant normalization algorithm resolving mismatched manufacturer model numbers and tier schemas with 95%+ comparison accuracy."),
                ("AI Review Sentiment Synthesis: ", "Developed an NLP review sentiment analysis pipeline aggregating thousands of verified buyer reviews into positive/negative consensus tags and 90-day price trend graphs.")
            ]
        },
        {
            "title": "WhatsApp Business AI Agent — Enterprise Messaging & Meeting Scheduler",
            "tech": "Python, Meta WhatsApp Cloud API, Google Calendar API v3, Webhooks",
            "link": "github.com/Ahmad-Haassan/WhatsApp-Business-Agent",
            "bullets": [
                ("Meta Cloud API Webhook Gateway: ", "Implemented a secure webhook receiver with SHA256 signature verification (X-Hub-Signature-256) to ingest real-time customer WhatsApp messages."),
                ("Autonomous Appointment Scheduling: ", "Integrated Google Calendar API v3 to parse conversational dates/times, verify consultant availability, reserve calendar slots, and issue automated meeting links."),
                ("Intent Classification Router: ", "Deployed an AI-driven intent router handling automated customer service FAQs, order inquiries, and escalation hand-offs with an average response latency under 2.5 seconds.")
            ]
        },
        {
            "title": "JARVIS AI — Offline Wake-Word Voice Assistant & Desktop Automation Suite",
            "tech": "Python 3.10+, Faster-Whisper, Vosk, PyQt6, Windows OS Automation",
            "link": "github.com/Ahmad-Haassan/JARVIS-Voice-Agent",
            "bullets": [
                ("Zero-Cloud Speech Pipeline: ", "Architected a 100% private, dual-stage speech engine combining an ultra-low CPU offline Vosk acoustic model for wake-word idling and local quantized Faster-Whisper for local STT."),
                ("System-Wide Macro Automation: ", "Programmed deep Windows OS automation indexing installed applications and executing multi-step macros (\"Start Work Mode\" to launch IDEs, terminals, and documentation tabs)."),
                ("Cyberpunk HUD Visualizer: ", "Built a hardware-accelerated PyQt6 desktop interface featuring dynamic pulsing audio waveform visualizers consuming under 3.5% idle CPU.")
            ]
        },
        {
            "title": "Developer Platform & Digital Command Center",
            "tech": "Next.js 16 (App Router), React 19, TypeScript, Tailwind CSS v4, Motion",
            "link": "github.com/Ahmad-Haassan/ahmad-portfolio",
            "bullets": [
                ("Interactive System Dashboard: ", "Engineered a developer telemetry portal showcasing active prototypes, verified skill trees, and coursework progression with fluid micro-interactions."),
                ("Accessibility & Keyboard Navigation: ", "Built a system-wide Command Palette (Cmd + K) for instant keyboard shortcuts across deep-dive case studies and contact modules.")
            ]
        }
    ]

    for proj in projects:
        # Title line
        header_p = doc.add_paragraph()
        header_p.paragraph_format.space_before = Pt(4)
        header_p.paragraph_format.space_after = Pt(1)
        header_p.paragraph_format.keep_with_next = True
        
        t_run = header_p.add_run(proj["title"])
        t_run.bold = True
        t_run.font.name = "Arial"
        t_run.font.size = Pt(9.5)
        t_run.font.color.rgb = COLOR_PRIMARY
        
        # Tech & Link line
        meta_p = doc.add_paragraph()
        meta_p.paragraph_format.space_before = Pt(0)
        meta_p.paragraph_format.space_after = Pt(2)
        meta_p.paragraph_format.keep_with_next = True
        
        tech_run = meta_p.add_run(proj["tech"] + "  |  " + proj["link"])
        tech_run.font.name = "Arial"
        tech_run.font.size = Pt(8.5)
        tech_run.font.italic = True
        tech_run.font.color.rgb = COLOR_MUTED
        
        # Bullets
        for bold_part, text_part in proj["bullets"]:
            bp = doc.add_paragraph(style="List Bullet")
            bp.paragraph_format.space_before = Pt(0)
            bp.paragraph_format.space_after = Pt(1.5)
            bp.paragraph_format.line_spacing = 1.15
            
            b_bold = bp.add_run(bold_part)
            b_bold.bold = True
            b_bold.font.name = "Arial"
            b_bold.font.size = Pt(9)
            b_bold.font.color.rgb = COLOR_PRIMARY
            
            b_text = bp.add_run(text_part)
            b_text.font.name = "Arial"
            b_text.font.size = Pt(9)
            b_text.font.color.rgb = COLOR_DARK

    # 5. Education
    add_heading_with_border("Education")
    
    edu_p = doc.add_paragraph()
    edu_p.paragraph_format.space_before = Pt(2)
    edu_p.paragraph_format.space_after = Pt(1)
    edu_p.paragraph_format.keep_with_next = True
    
    deg_run = edu_p.add_run("Bachelor of Science in Computer Science (BSCS)")
    deg_run.bold = True
    deg_run.font.name = "Arial"
    deg_run.font.size = Pt(9.5)
    deg_run.font.color.rgb = COLOR_PRIMARY
    
    grad_run = edu_p.add_run("  |  Expected Graduation: 2026")
    grad_run.font.name = "Arial"
    grad_run.font.size = Pt(9)
    grad_run.font.color.rgb = COLOR_MUTED
    
    course_p = doc.add_paragraph()
    course_p.paragraph_format.space_before = Pt(0)
    course_p.paragraph_format.space_after = Pt(4)
    course_p.paragraph_format.line_spacing = 1.15
    
    c_lbl = course_p.add_run("Relevant Academic Coursework: ")
    c_lbl.bold = True
    c_lbl.font.name = "Arial"
    c_lbl.font.size = Pt(8.5)
    c_lbl.font.color.rgb = COLOR_MUTED
    
    c_val = course_p.add_run(
        "Data Structures & Algorithms, Object-Oriented Programming (C++), Operating Systems & Concurrency, "
        "Database Management Systems (SQL), Computer Networks, Software Engineering Principles, Artificial Intelligence."
    )
    c_val.font.name = "Arial"
    c_val.font.size = Pt(8.5)
    c_val.font.color.rgb = COLOR_DARK

    # Save to both public/ and workspace root
    doc.save("public/Ahmad_Hassan_Resume.docx")
    doc.save("Ahmad_Hassan_Resume.docx")
    print("Successfully generated Ahmad_Hassan_Resume.docx in public/ and root!")

if __name__ == "__main__":
    create_resume()
